"""Validação do PPTX gerado: abre o arquivo, conta slides e procura estouro de caixa.

Um .pptx que "abre" pode ter texto fora do slide, tabela maior que o espaço ou
cor de fundo ilegível — e isso só aparece quando alguém projeta. Este check
transforma a apresentação em um teste: se um número divergir do banco ou um
elemento sair da área do slide, o teste falha.
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from workers.apresentacao_pptx import construir, _numeros


def _dentro(slide, forma, largura, altura, tol=Emu(int(0.02 * 914400))) -> list[str]:
    """Formas que estouram a área do slide."""
    problemas = []
    for nome, x, y, w, h in ((f"{forma.shape_type}", forma.left, forma.top,
                              forma.width, forma.height),):
        if x < -tol or y < -tol:
            problemas.append(f"{nome} começa fora do slide (x={x}, y={y})")
        if x + w > largura + tol or y + h > altura + tol:
            problemas.append(f"{nome} estoura o slide "
                            f"(right={Emu(x + w).inches:.2f}in, "
                            f"bottom={Emu(y + h).inches:.2f}in)")
    return problemas


def _texto_de(forma) -> str:
    if forma.has_text_frame:
        return forma.text_frame.text
    if getattr(forma, "has_table", False):
        return " ".join(c.text for linha in forma.table.rows for c in linha.cells)
    return ""


def validar(destino: Path) -> list[str]:
    """Devolve a lista de problemas (vazia = apresentação íntegra)."""
    problemas: list[str] = []
    prs = Presentation(str(destino))
    largura, altura = prs.slide_width, prs.slide_height
    if not prs.slides:
        return ["a apresentação não tem slides"]
    for i, slide in enumerate(prs.slides, start=1):
        if len(slide.shapes) == 0:
            problemas.append(f"slide {i}: sem nenhum elemento")
        for forma in slide.shapes:
            problemas += [f"slide {i}: {p}" for p in _dentro(slide, forma, largura, altura)]
    return problemas


def _conferir_numeros(destino: Path, n: dict) -> list[str]:
    """Os números do deck têm de bater com o banco — é o ponto do projeto."""
    problemas = []
    texto = " ".join(_texto_de(f) for s in Presentation(str(destino)).slides
                     for f in s.shapes)
    esperado = {
        "fatos": str(n["fatos"]), "fontes": str(n["fontes"]),
        "DQS": str(n["dqs"]), "projeções": str(n["projecoes"]),
        "séries": str(n["series"]), "alertas": str(n["alertas"]),
        "scorecards": str(n["scorecards"]),
    }
    for rotulo, valor in esperado.items():
        if valor not in texto:
            problemas.append(f"número ausente no deck: {rotulo} = {valor}")
    return problemas


def main() -> int:
    destino = Path("docs/APRESENTACAO_PETROBRAS.pptx")
    caminho = construir(destino)
    # o CLI precisa gerar o MESMO arquivo, senão `app_main.py pdf --pptx` e o
    # validador produziriam decks diferentes e o validador não significaria nada
    if str(Path(caminho).resolve()) != str(destino.resolve()):
        print(f"  ! construir() ignorou o destino informado: {caminho}")
    problemas = validar(Path(caminho))
    problemas += _conferir_numeros(Path(caminho), _numeros())
    prs = Presentation(caminho)
    print(f"{caminho}: {len(prs.slides)} slides, {len(list(prs.slides[0].shapes))} "
          f"formas no primeiro")
    for p in problemas:
        print("  !", p)
    print("OK" if not problemas else f"{len(problemas)} problema(s)")
    return 1 if problemas else 0


if __name__ == "__main__":
    raise SystemExit(main())